#!/usr/bin/env python3
"""Группы игроков, рассылка по группе и повторяющиеся письма.

    python3 tests/test_groups.py

База временная и своя. Наружу ничего не уходит: бот подменён заглушкой, а
отдельная проверка следит за главным обещанием раздела — письмо уходит
ТОЛЬКО в личку и только после подтверждения. Случайная рассылка на всю
команду с одного нажатия — ровно то, чего здесь быть не должно.
"""

from __future__ import annotations

import asyncio
import os
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, List

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from fake_tg import (FakeBot, FakeContext, FakeMessage, FakeQuery, FakeUpdate,
                     FakeUser, buttons_of)

TMP = Path(tempfile.mkdtemp(prefix="groups-test-")) / "bot.db"

COACH = FakeUser(uid=800100, username="coach")
BOT = FakeBot()

bad: List[str] = []


def check(cond: bool, what: str) -> None:
    print(("  ✅ " if cond else "  ❌ ") + what)
    if not cond:
        bad.append(what)


def setup() -> Any:
    os.environ.setdefault("BOT_TOKEN", "0:test")
    os.environ["ADMIN_USER_IDS"] = str(COACH.id)
    os.environ.setdefault("DAEMON_LOG_PATH", str(ROOT / "tests" / "test.log"))
    os.environ["GOOGLE_SHEETS_CREDENTIALS"] = ""
    os.environ["SPREADSHEET_ID"] = ""
    import sheets_cache
    sheets_cache.DB_PATH = TMP
    sheets_cache.init_db()
    now = sheets_cache.now_iso()
    with sheets_cache.get_connection() as conn:
        conn.execute("DELETE FROM players")
        conn.execute("DELETE FROM player_links")
        conn.execute("DELETE FROM league_teams")
        # У одного даты рождения нет намеренно: заявку с пустой клеткой
        # вернут, и выгрузка обязана предупредить об этом заранее.
        people = [("Иванов", "Иван", "2001-09-22"),
                  ("Петров", "Пётр", "1998-03-04"),
                  ("Сидоров", "Семён", "")]
        for i, (sur, name, born) in enumerate(people, start=2):
            conn.execute(
                "INSERT INTO players (row_index, surname, name, birthday, "
                "synced_at) VALUES (?, ?, ?, ?, ?)", (i, sur, name, born, now))
        # Двое запускали бота, третий — нет: рассылка обязана это различать.
        for row, uid in ((2, 900002), (3, 900003)):
            conn.execute(
                "INSERT INTO player_links (tg_user_id, username, player_row, "
                "linked_at) VALUES (?, '', ?, ?)", (str(uid), row, now))
        conn.execute(
            "INSERT INTO league_teams (source, team_id, name, league, ours, "
            "fetched_at) VALUES ('infobasket','36502','PULL UP',"
            "'Летняя лига · Группа 4',1,?)", (now,))
        conn.commit()
    import bot_daemon as bd
    bd._get_spreadsheet = lambda: None
    return bd


async def press(bd, data: str, who: FakeUser = COACH):
    q = FakeQuery(data, who, BOT)
    await bd.handle_group_callback(FakeUpdate(query=q, user=who), FakeContext(BOT))
    last = (q.screens or [{"text": "", "markup": None}])[-1]
    return last["text"], last["markup"]


async def say(bd, text: str, who: FakeUser = COACH):
    msg = FakeMessage(text=text, bot=BOT, user=who)
    try:
        await bd.handle_group_text(FakeUpdate(message=msg, user=who),
                                   FakeContext(BOT))
    except Exception as exc:                     # ApplicationHandlerStop — норма
        if type(exc).__name__ != "ApplicationHandlerStop":
            raise
    last = (msg.replies or [{"text": "", "markup": None}])[-1]
    return last["text"], last["markup"]


def cbs(markup) -> List[str]:
    return [b.callback_data for b in buttons_of(markup) if b.callback_data]


def find(markup, needle: str) -> str:
    for b in buttons_of(markup):
        if needle.lower() in (b.text or "").lower():
            return b.callback_data
    return ""


# ─────────────────────────── сценарии ──────────────────────────────────────


async def test_create_and_fill(bd) -> int:
    """Завести группу, назвать вольно, набрать состав, привязать лигу."""
    print("\n=== группа заводится и наполняется ===")
    import player_groups as pg

    text, markup = await press(bd, "pg:main")
    check("Групп" in text, "раздел открылся")
    check("pg:new" in cbs(markup), "есть кнопка создания")

    await press(bd, "pg:new")
    text, _ = await say(bd, "Второй состав")
    check("Второй состав" in text, f"имя вольное: {text.splitlines()[0]}")
    gid = pg.groups()[0]["id"]

    # Повтор имени не заводит вторую такую же группу — и в другом регистре тоже.
    await press(bd, "pg:new")
    text, _ = await say(bd, "второй СОСТАВ")
    check("уже есть" in text, "дубль имени отбит независимо от регистра")
    check(len(pg.groups()) == 1, "групп по-прежнему одна")

    _, markup = await press(bd, f"pg:who:{gid}:0")
    picks = [c for c in cbs(markup) if c.startswith(f"pg:t:{gid}:")]
    check(len(picks) == 3, f"в списке все игроки: {len(picks)}")
    for c in picks:
        await press(bd, c)
    check(len(pg.member_rows(gid)) == 3, "все трое в группе")

    # Повторное нажатие убирает — список работает как переключатель.
    await press(bd, picks[0])
    check(len(pg.member_rows(gid)) == 2, "нажал ещё раз — вышел из группы")
    await press(bd, picks[0])

    _, markup = await press(bd, f"pg:lg:{gid}")
    where = [c for c in cbs(markup) if c.startswith(f"pg:lgset:{gid}:")]
    await press(bd, where[0])
    text, _ = await press(bd, f"pg:g:{gid}")
    check("Летняя лига" in text, f"лига привязана: {text}")

    await press(bd, f"pg:lgset:{gid}:-1")
    check(not pg.group(gid)["league_source"], "и отвязывается обратно")
    await press(bd, where[0])
    return gid


async def test_broadcast_asks_first(bd, gid: int) -> None:
    """Письмо уходит только после подтверждения и только в личку."""
    print("\n=== рассылка спрашивает, прежде чем уйти ===")
    before = len(BOT.sent)

    await press(bd, f"pg:free:{gid}")
    text, markup = await say(bd, "Завтра зал в 19:00, не опаздывайте")
    check(len(BOT.sent) == before, "после набора текста НИЧЕГО не ушло")
    check("Завтра зал" in text, "письмо показано целиком")
    check(f"pg:go:{gid}" in cbs(markup), "есть кнопка отправки")

    report, _ = await press(bd, f"pg:go:{gid}")
    fresh = BOT.sent[before:]
    check(len(fresh) == 2, f"ушло ровно двоим из троих: {len(fresh)}")
    check(all(str(m.get("chat_id")) in ("900002", "900003") for m in fresh),
          "и обоим в личку, по их собственным id")
    check(all("Завтра зал" in str(m.get("text", "")) for m in fresh),
          "текст дошёл без изменений")

    check("Отправлено: 2" in report, f"тренеру отчитались: {report.splitlines()[0]}")
    check("не запускали бота: 1" in report.lower(),
          "и про недоступного напомнили в отбивке")


async def test_silent_are_named(bd, gid: int) -> None:
    """Тот, кто не запускал бота, назван — а не потерян молча."""
    print("\n=== молчуны названы ===")
    text, _ = await press(bd, f"pg:g:{gid}")
    check("Сидоров" in text, f"третий назван как недоступный: {text}")
    text, _ = await press(bd, f"pg:send:{gid}")
    check("не дойдёт" in text.lower(), "и на экране письма предупредили")


async def test_templates(bd, gid: int) -> None:
    """Шаблон сохраняется, рассылается и переписывается."""
    print("\n=== шаблоны писем ===")
    import player_groups as pg

    await press(bd, "pg:tnew")
    await say(bd, "Зал в четверг")
    text, _ = await say(bd, "Четверг, зал на Ленина, 20:00")
    check("сохранён" in text.lower(), f"шаблон записан: {text.splitlines()[0]}")
    tid = pg.templates()[0]["id"]

    before = len(BOT.sent)
    _, markup = await press(bd, f"pg:tto:{tid}")
    to_group = [c for c in cbs(markup) if c.startswith(f"pg:use2:{tid}:")]
    check(bool(to_group), "предложили выбрать группу")
    text, markup = await press(bd, to_group[0])
    check(len(BOT.sent) == before, "выбор группы сам по себе ничего не шлёт")
    check("Ленина" in text, "показали текст шаблона")
    await press(bd, f"pg:go:{gid}")
    check(len(BOT.sent) - before == 2, "после подтверждения ушло двоим")

    await press(bd, f"pg:ted:{tid}")
    await say(bd, "Четверг, зал на Мира, 20:30")
    check("Мира" in pg.template(tid)["body"], "текст шаблона переписан")
    check(len(pg.templates()) == 1, "второй шаблон не завёлся")


async def test_repeat(bd, gid: int) -> None:
    """Повтор: день, время, и он уходит сам — но ровно один раз в день."""
    print("\n=== повторяющееся письмо ===")
    import player_groups as pg

    await press(bd, f"pg:rfree:{gid}")
    await say(bd, "Напоминаю про зал")
    # Понедельник, 19:30
    await press(bd, f"pg:rtime:{gid}:0:0")
    text, _ = await say(bd, "19:30")
    check("понедельник" in text.lower(), f"повтор заведён: {text.splitlines()[0]}")

    reps = pg.repeats(gid)
    check(len(reps) == 1, "повтор один")
    rid = reps[0]["id"]

    monday = datetime(2026, 8, 24, 19, 35)       # понедельник, чуть позже срока
    check(bool(pg.due(monday)), "в свой день и час — пора")
    check(not pg.due(datetime(2026, 8, 25, 19, 35)), "во вторник не пора")
    check(not pg.due(datetime(2026, 8, 24, 9, 0)), "утром ещё рано")
    check(not pg.due(datetime(2026, 8, 24, 23, 0)),
          "и через четыре часа уже поздно — вчерашнее письмо не догоняем")

    before = len(BOT.sent)
    pg.mark_sent(rid, monday.date())
    check(not pg.due(monday), "отправленное сегодня второй раз не уходит")

    pg.repeat_switch(rid)
    check(not pg.repeats(gid)[0]["active"], "повтор выключается")
    check(not pg.due(datetime(2026, 8, 31, 19, 35)), "выключенный молчит")
    pg.repeat_switch(rid)
    check(len(BOT.sent) == before, "вся эта возня никому ничего не отправила")


async def test_delete_takes_repeats(bd) -> None:
    """Удаление группы уносит её повторы — иначе письмо ушло бы в никуда."""
    print("\n=== удаление группы ===")
    import player_groups as pg
    gid, _ = pg.create("На выброс")
    pg.repeat_add(gid, 2, "18:00", 0, "текст")
    check(len(pg.repeats(gid)) == 1, "у группы есть повтор")
    pg.delete(gid)
    check(all(int(r["group_id"]) != int(gid) for r in pg.repeats()),
          "повтор ушёл вместе с группой")


async def test_nothing_leaks_to_chat(bd) -> None:
    """Ни одно сообщение раздела не ушло в групповой чат."""
    print("\n=== в общий чат не уходит ничего ===")
    group_chats = [m for m in BOT.sent
                   if str(m.get("chat_id", "")).startswith("-")]
    check(not group_chats, f"сообщений в группы: {len(group_chats)}")
    threads = [m for m in BOT.sent if m.get("message_thread_id")]
    check(not threads, f"сообщений в топики: {len(threads)}")


async def test_clear_pending_knows_us(bd) -> None:
    """/start обязан закрывать и диалоги этого раздела."""
    print("\n=== /start закрывает диалоги групп ===")
    bd._awaiting_group[COACH.id] = "new"
    bd._group_letter[COACH.id] = {"gid": 1, "body": "черновик"}
    bd._clear_pending(COACH.id)
    check(COACH.id not in bd._awaiting_group, "ожидание ввода снято")
    check(COACH.id not in bd._group_letter, "черновик письма забыт")


async def test_export_for_the_application(bd, gid: int) -> None:
    """Выгрузка состава в таблицу — под заявку в лигу."""
    print("\n=== выгрузка для заявки ===")
    import roster_export

    _, markup = await press(bd, f"pg:g:{gid}")
    check(f"pg:csv:{gid}" in cbs(markup), "кнопка есть в группе")
    _, markup = await press(bd, "coach:main")
    team = bd._team_markup().inline_keyboard
    check("coach:csv" in [b.callback_data for row in team for b in row],
          "и в разделе «Команда»")

    before = len(BOT.docs) if hasattr(BOT, "docs") else None
    msg = FakeMessage(text="", bot=BOT, user=COACH)
    q = FakeQuery(f"pg:csv:{gid}", COACH, BOT)
    q.message = msg
    await bd.handle_group_callback(FakeUpdate(query=q, user=COACH), FakeContext(BOT))
    sent = [d for d in getattr(msg, "documents", [])]
    check(len(sent) == 1, f"файл отправлен: {len(sent)}")
    if not sent:
        return
    raw = sent[0]["document"].getvalue()
    name = sent[0]["filename"]
    check(name.startswith("Заявка"), f"имя файла говорит, что это: {name}")
    check("Без даты рождения" in sent[0]["caption"],
          "про пустую дату рождения предупредили: "
          + sent[0]["caption"].splitlines()[-1][:60])

    people = pg_members_count(gid)
    if name.endswith(".xlsx"):
        check(raw[:2] == b"PK", "это настоящая книга Excel, а не переименованный csv")
        import io
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(raw)).active
        table = list(sheet.iter_rows(values_only=True))
        check(table[0][:3] == ("№", "Фамилия", "Имя"), f"шапка на месте: {table[0]}")
        check(len(table) == 1 + len(people),
              f"строк по числу людей в группе: {len(table) - 1}")
        check(isinstance(table[1][0], int), "номер лежит числом, а не текстом")
        # Ищем столбец по названию: их состав менялся, и жёсткий индекс
        # молча начал бы проверять соседнюю колонку.
        at = list(table[0]).index("Дата рождения")
        born = [r[at] for r in table[1:] if r[at] is not None]
        # Дата — ТЕКСТОМ. Настоящая дата открывается верно в Excel, но
        # просмотр Telegram рисует её по-своему: 22.09.2001 показывалось как
        # «265.09.2001» — это день ГОДА. Заявку чаще всего смотрят прямо в
        # чате, и врать при первом взгляде она не должна.
        check(born and all(isinstance(b, str) for b in born),
              f"дата рождения лежит текстом: {born[:2]}")
        check(any("22.09.2001" == b for b in born),
              f"и ровно в том виде, в каком её читают: {born[:2]}")
        check(sheet.freeze_panes == "A2", "шапка не уезжает при прокрутке")
    else:
        body = raw.decode("utf-8")
        check(body.startswith("\ufeff"), "в начале BOM — иначе Excel даст кракозябры")
        head = body.splitlines()[0]
        check(head.count(";") == 5 and head.startswith("\ufeff№;Фамилия;Имя"),
              f"шапка через точку с запятой: {head}")
        check("22.09.2001" in body, "дата рождения по-русски")
        check("2001-09-22" not in body, "и ISO-вида в файле не осталось")
        lines = [l for l in body.splitlines() if l.strip()]
        check(len(lines) == 1 + len(people),
              f"строк по числу людей в группе: {len(lines) - 1}")


async def test_numbers_come_from_protocols(bd, gid: int) -> None:
    """Номер доносится из протоколов, если лига его не отдала.

    Заявка SLPRO номеров не содержит вовсе — в ответе только id, имя и дата
    рождения. Но в каждом протоколе записано, под каким номером человек
    выходил, и эти протоколы бот и так качает."""
    print("\n=== номер из протокола ===")
    import league_sync
    import sheets_cache

    now = sheets_cache.now_iso()
    with sheets_cache.get_connection() as conn:
        conn.execute("DELETE FROM game_player_stats WHERE player_id IN ('555','556')")
        # Один и тот же человек в двух играх под разными номерами.
        for game, day, num in (("g1", "2026-07-01", "8"), ("g2", "2026-08-01", "12")):
            conn.execute(
                "INSERT INTO game_player_stats (source, game_id, player_id, "
                "team_id, number, game_date, fetched_at) VALUES ('slpro', ?, "
                "'555', '707', ?, ?, ?)", (game, num, day, now))
        conn.commit()

    got = league_sync._numbers_from_protocols("slpro", ["555", "556"])
    check(got.get("555") == "12",
          f"взят номер из самого свежего протокола: {got.get('555')}")
    check("556" not in got, "кого в протоколах нет — без номера")
    check(league_sync._numbers_from_protocols("slpro", []) == {},
          "пустой запрос не ходит в базу зря")

    with sheets_cache.get_connection() as conn:
        conn.execute("DELETE FROM game_player_stats WHERE player_id IN ('555','556')")
        conn.commit()


async def test_shirt_numbers_in_export(bd, gid: int) -> None:
    """Игровой номер в выгрузке: из заявки лиги, а не из листа.

    В листе «Игроки» номера нет — его ведёт лига. Берём по привязанному
    профилю, а кого не привязали — по имени из реестра лиг."""
    print("\n=== игровые номера в заявке ===")
    import player_names
    import roster_export
    import sheets_cache

    now = sheets_cache.now_iso()
    with sheets_cache.get_connection() as conn:
        conn.execute("DELETE FROM league_rosters")
        conn.execute("DELETE FROM player_identities")
        for pid, num in (("111", "7"), ("222", "13"), ("333", "9")):
            conn.execute(
                "INSERT INTO league_rosters (source, team_id, player_id, "
                "number, fetched_at) VALUES ('infobasket','36502',?,?,?)",
                (pid, num, now))
        # Иванов привязан профилем — точный путь.
        conn.execute(
            "INSERT INTO player_identities (tg_user_id, source, player_id, "
            "linked_at) VALUES ('900002','infobasket','111',?)", (now,))
        conn.commit()
    # Петров не привязан: его находим по имени из реестра лиг.
    player_names.put("infobasket", "222", "Петров Пётр")
    player_names.put("infobasket", "333", "Кто-то Чужой")

    numbers = roster_export.numbers_by_row()
    check(numbers.get(2) == "7", f"привязанный — по профилю: {numbers.get(2)}")
    check(numbers.get(3) == "13", f"непривязанный — по имени: {numbers.get(3)}")
    check(4 not in numbers, "кого в заявке лиги нет — без номера")

    people = roster_export.group_people(gid)
    table = roster_export.rows(people)
    check(table[0][3] == "Игровой номер", f"столбец на месте: {table[0]}")
    shirts = {line[1]: line[3] for line in table[1:]}
    check(shirts.get("Иванов") == "7", f"номер попал в строку: {shirts}")
    check(shirts.get("Сидоров") == "", "чужой номер никому не приписали")

    with sheets_cache.get_connection() as conn:
        conn.execute("DELETE FROM league_rosters")
        conn.execute("DELETE FROM player_identities")
        conn.commit()
    check(not roster_export.numbers_by_row(),
          "без заявок лиг номеров нет, и выгрузка не падает")


async def test_csv_still_works_without_openpyxl(bd, gid: int) -> None:
    """Без библиотеки выгрузка не падает, а отдаёт CSV.

    Библиотеку ставят руками в окружение бота, и на новом сервере её может не
    оказаться. Заявка нужна тренеру в любом случае — пусть таблицей попроще."""
    print("\n=== без openpyxl отдаём CSV ===")
    import roster_export
    was = roster_export.has_xlsx
    roster_export.has_xlsx = lambda: False
    try:
        people = roster_export.group_people(gid)
        data, name = roster_export.build(people, "Второй состав")
    finally:
        roster_export.has_xlsx = was
    check(name.endswith(".csv"), f"формат сменился: {name}")
    body = data.decode("utf-8")
    check(body.startswith("\ufeff"), "BOM на месте — Excel прочтёт кириллицу")
    check("22.09.2001" in body, "дата рождения по-русски и здесь")
    check(len([l for l in body.splitlines() if l.strip()]) == 1 + len(people),
          "и люди все на месте")


def pg_members_count(gid: int):
    import player_groups as pg
    return pg.members(gid)


async def run() -> None:
    bd = setup()
    gid = await test_create_and_fill(bd)
    await test_broadcast_asks_first(bd, gid)
    await test_silent_are_named(bd, gid)
    await test_templates(bd, gid)
    await test_repeat(bd, gid)
    await test_export_for_the_application(bd, gid)
    await test_numbers_come_from_protocols(bd, gid)
    await test_shirt_numbers_in_export(bd, gid)
    await test_csv_still_works_without_openpyxl(bd, gid)
    await test_delete_takes_repeats(bd)
    await test_nothing_leaks_to_chat(bd)
    await test_clear_pending_knows_us(bd)


def main() -> int:
    print(f"База: {TMP}")
    asyncio.run(run())
    print("\n" + "=" * 60)
    if bad:
        print(f"НЕ ПРОШЛО ({len(bad)}):")
        for b in bad:
            print("  • " + b)
        return 1
    print("ГРУППЫ И РАССЫЛКИ: ВСЁ ЗЕЛЁНОЕ")
    return 0


if __name__ == "__main__":
    sys.exit(main())
